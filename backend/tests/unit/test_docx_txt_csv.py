"""Unit tests for direct format parsing (DOCX, TXT, CSV)."""

import io
import pytest
from docx import Document

from app.services.pipeline.docx_txt_csv import DirectFormatParser


@pytest.mark.asyncio
async def test_parse_txt() -> None:
    """Verify parsing of plain text files."""
    parser = DirectFormatParser()
    txt_bytes = "Plain text document content.\nSecond line here.".encode("utf-8")
    pages = await parser.parse(txt_bytes, "txt")

    assert len(pages) == 1
    assert "Plain text document content." in pages[0].raw_text
    assert not pages[0].is_scanned


@pytest.mark.asyncio
async def test_parse_csv() -> None:
    """Verify parsing of CSV files into structured Markdown tables."""
    parser = DirectFormatParser()
    csv_bytes = b"Name,Age,Role\nAlice,30,Engineer\nBob,35,Manager\n"
    pages = await parser.parse(csv_bytes, "csv")

    assert len(pages) == 1
    assert "Alice" in pages[0].raw_text
    assert "Engineer" in pages[0].raw_text
    assert "|" in pages[0].raw_text


@pytest.mark.asyncio
async def test_parse_docx() -> None:
    """Verify parsing of DOCX files."""
    parser = DirectFormatParser()
    doc = Document()
    doc.add_paragraph("First paragraph of docx.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Header 1"
    table.rows[0].cells[1].text = "Header 2"

    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    pages = await parser.parse(docx_bytes, "docx")
    assert len(pages) == 1
    assert "First paragraph of docx." in pages[0].raw_text
    assert "Header 1" in pages[0].raw_text
