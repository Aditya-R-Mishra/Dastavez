"""Direct text and structured table parsers for DOCX, TXT, and CSV files."""

import io
from typing import List
from starlette.concurrency import run_in_threadpool

from app.core.exceptions import InvalidFileTypeError
from app.services.pipeline.text_extraction import ExtractedPage


class DirectFormatParser:
    """Extracts text content from non-PDF direct formats (DOCX, TXT, CSV)."""

    def _sync_parse_docx(self, file_bytes: bytes) -> List[ExtractedPage]:
        """Parse text paragraphs and tables from Microsoft Word (.docx) files."""
        from docx import Document

        doc = Document(io.BytesIO(file_bytes))
        content_parts = []

        # Extract text paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                content_parts.append(text)

        # Extract table rows
        for table in doc.tables:
            table_lines = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_lines.append(" | ".join(cells))
            if table_lines:
                content_parts.append("\n".join(table_lines))

        full_text = "\n\n".join(content_parts)
        return [ExtractedPage(page_number=1, raw_text=full_text, is_scanned=False)]

    def _sync_parse_txt(self, file_bytes: bytes) -> List[ExtractedPage]:
        """Decode plain text (.txt) files using UTF-8 with fallback."""
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="replace")
        return [ExtractedPage(page_number=1, raw_text=text.strip(), is_scanned=False)]

    def _sync_parse_csv(self, file_bytes: bytes) -> List[ExtractedPage]:
        """Convert CSV tabular data into structured Markdown table format."""
        import pandas as pd

        df = pd.read_csv(io.BytesIO(file_bytes))
        markdown_table = df.to_markdown(index=False)
        return [ExtractedPage(page_number=1, raw_text=markdown_table, is_scanned=False)]

    def _sync_parse(self, file_bytes: bytes, file_type: str) -> List[ExtractedPage]:
        clean_type = file_type.lower().strip(".").strip()
        if clean_type == "docx":
            return self._sync_parse_docx(file_bytes)
        elif clean_type == "txt":
            return self._sync_parse_txt(file_bytes)
        elif clean_type == "csv":
            return self._sync_parse_csv(file_bytes)
        else:
            raise InvalidFileTypeError(file_type, ["docx", "txt", "csv", "pdf", "png", "jpg", "jpeg"])

    async def parse(self, file_bytes: bytes, file_type: str) -> List[ExtractedPage]:
        """Asynchronously parse content from DOCX, TXT, or CSV file bytes."""
        return await run_in_threadpool(self._sync_parse, file_bytes, file_type)
